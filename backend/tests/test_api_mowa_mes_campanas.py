"""Pruebas de los endpoints de campanas y reportes de MOWA MES (corte B6b).

Sustituyen las dependencias por los servicios reales sobre repositorios en
memoria; el archivo se escribe con el adaptador .xlsx de verdad y el reporte se
lee con el lector de verdad. Datos sinteticos.
"""

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.adapters.input.lector_reporte_mowa_mes import LectorReporteMowaMes
from app.adapters.output.plataformas.mowa_mes.archivo_carga import escribir_archivo_carga
from app.api import contrato
from app.api.mowa_mes_campanas import (
    CABECERAS_ARCHIVO,
    obtener_servicio_campanas,
    obtener_servicio_reportes,
)
from app.core.entities.gestiones_digitales import ConfiguracionSupervision
from app.core.entities.mowa_mes_reporte import COLUMNAS_REPORTE
from app.core.services.plataformas.mowa_mes.campana import CampanasMowaMesService
from app.core.services.plataformas.mowa_mes.reportes import ReportesMowaMesService
from app.core.services.seleccion_cartera.servicio import ConsultaCarteraService
from app.main import app
from tests.test_gestiones_digitales import PROCEDENCIAS, SUPERVISORES, RepositorioSupervisionFalso
from tests.test_mowa_mes_campana import (
    HUELLA_ORIGINAL,
    MOMENTO,
    RepositorioCampanasMemoria,
    RepositorioCarteraFalso,
    producto,
)
from tests.test_mowa_mes_speech import WHATSAPP, RepositorioMowaMesFalso

CUERPO = {
    "fecha_corte": "2026-09-10",
    "cantidad": 100,
    "programacion": "hora_determinada",
    "envios": ["2026-09-14T09:00:00"],
}


class Entorno:
    def __init__(self, productos=None, supervisores=SUPERVISORES, cargados_mes=0, vigente=True):
        self.cartera = RepositorioCarteraFalso(
            productos if productos is not None else [producto(1), producto(2, telefono=None)],
            vigente,
        )
        self.mowa_mes = RepositorioMowaMesFalso(WHATSAPP)
        self.supervision = RepositorioSupervisionFalso()
        self.supervision.configuracion = ConfiguracionSupervision(PROCEDENCIAS, tuple(supervisores))
        self.campanas = RepositorioCampanasMemoria(cargados_mes)

    def campanas_servicio(self):
        return CampanasMowaMesService(
            ConsultaCarteraService(self.cartera),
            self.mowa_mes,
            self.supervision,
            self.campanas,
            escribir_archivo_carga,
            reloj=lambda: MOMENTO,
        )

    def reportes_servicio(self):
        return ReportesMowaMesService(self.campanas, LectorReporteMowaMes())


def _cliente(entorno: Entorno) -> TestClient:
    app.dependency_overrides[obtener_servicio_campanas] = entorno.campanas_servicio
    app.dependency_overrides[obtener_servicio_reportes] = entorno.reportes_servicio
    return TestClient(app)


@pytest.fixture(autouse=True)
def _limpiar_dependencias():
    yield
    app.dependency_overrides.clear()


def _crear(client: TestClient, **extra):
    return client.post(
        "/mowa-mes/campanas", json={**CUERPO, "speech_huella": HUELLA_ORIGINAL, **extra}
    )


# --- Previsualizacion ---------------------------------------------------


def test_previsualizacion_de_una_campana_valida() -> None:
    client = _cliente(Entorno())

    respuesta = client.post("/mowa-mes/campanas/previsualizacion", json=CUERPO)

    cuerpo = respuesta.json()
    assert respuesta.status_code == 200
    assert cuerpo["puede_crear"] is True and cuerpo["errores"] == []
    assert cuerpo["speech"] == {"id": 1, "nombre": "Speech original", "huella": HUELLA_ORIGINAL}
    assert (cuerpo["productos_cargados"], cuerpo["supervision_cargados"]) == (1, 5)
    assert cuerpo["exclusiones_por_codigo"] == [
        {"codigo": "telefono_invalido", "tipo": "exclusion", "cantidad": 1}
    ]
    assert cuerpo["exclusiones"]["exclusiones"] == [
        {"pagare": producto(2)["pagare"], "codigo": "telefono_invalido"}
    ]
    assert [m["supervision"] for m in cuerpo["muestra"]] == [True] * 5 + [False]
    assert cuerpo["muestra"][5]["largo"] == len(cuerpo["muestra"][5]["mensaje"])
    assert cuerpo["archivos_previstos_por_filas"] == [{"numero": 1, "filas": 6, "supervision": 5}]
    assert cuerpo["limite"]["mes"] == "2026-09" and cuerpo["limite"]["excedido"] is False
    assert cuerpo["fecha_envio"] == "2026-09-14"
    assert {s["segmento"]: s["cantidad"] for s in cuerpo["productos_por_segmento"]}["1_a_8"] == 1
    assert cuerpo["descripcion"] == cuerpo["descripcion_sugerida"] == "CajaCusco"


def test_los_errores_de_campana_se_listan_sin_responder_error() -> None:
    client = _cliente(Entorno(supervisores=()))

    cuerpo = client.post("/mowa-mes/campanas/previsualizacion", json=CUERPO).json()

    assert cuerpo["puede_crear"] is False
    assert [(e["codigo"], e["tipo"]) for e in cuerpo["errores"]] == [("sin_supervisores", "error")]


def test_la_advertencia_de_limite_mensual() -> None:
    client = _cliente(Entorno(cargados_mes=2_500_000))

    cuerpo = client.post("/mowa-mes/campanas/previsualizacion", json=CUERPO).json()

    assert [(a["codigo"], a["tipo"]) for a in cuerpo["advertencias"]] == [
        ("limite_mensual_excedido", "advertencia")
    ]
    assert cuerpo["puede_crear"] is True


@pytest.mark.parametrize(
    ("entorno", "cambio", "estado"),
    [
        ({"vigente": False}, {}, 404),
        ({}, {"tipo_carga": "personalizada"}, 400),
        ({}, {"filtros": ["campo_inventado:igual:x"]}, 400),
        ({}, {"speech_id": 99}, 404),
        ({}, {"cantidad": 0}, 422),
        ({}, {"salida": "otra"}, 422),
    ],
)
def test_peticiones_de_previsualizacion_rechazadas(entorno, cambio, estado) -> None:
    client = _cliente(Entorno(**entorno))

    respuesta = client.post("/mowa-mes/campanas/previsualizacion", json={**CUERPO, **cambio})

    assert respuesta.status_code == estado


# --- Creacion y consulta ------------------------------------------------


def test_crear_y_consultar_una_campana() -> None:
    client = _cliente(Entorno())

    creada = _crear(client, descripcion="Campana sintetica")
    campana_id = creada.json()["id"]
    leida = client.get(f"/mowa-mes/campanas/{campana_id}")
    lista = client.get("/mowa-mes/campanas").json()
    exclusiones = client.get(
        f"/mowa-mes/campanas/{campana_id}/exclusiones", params={"codigo": "telefono_invalido"}
    ).json()
    otras = client.get(
        f"/mowa-mes/campanas/{campana_id}/exclusiones", params={"codigo": "sin_speech"}
    ).json()

    assert creada.status_code == 201
    cuerpo = leida.json()
    assert (cuerpo["descripcion"], cuerpo["total_cargados"], cuerpo["mes_imputacion"]) == (
        "Campana sintetica",
        6,
        "2026-09",
    )
    assert cuerpo["archivos"] == [
        {"numero": 1, "filas": 6, "supervision": 5, "bytes": cuerpo["archivos"][0]["bytes"]}
    ]
    assert cuerpo["speech"] == {"id": 1, "nombre": "Speech original"}
    # En el orden de la lista (entrelazada a proposito), cada uno con su DNI por procedencia.
    assert [s["documento"] for s in cuerpo["supervisores"]] == [
        "00000001",
        "00000004",
        "00000002",
        "00000005",
        "00000003",
    ]
    assert (lista["total"], lista["campanas"][0]["id"]) == (1, campana_id)
    assert exclusiones["total"] == 1 and otras["total"] == 0
    assert client.get("/mowa-mes/campanas/99").status_code == 404
    assert client.get("/mowa-mes/campanas/99/exclusiones").status_code == 404


@pytest.mark.parametrize(
    ("entorno", "extra", "estado"),
    [
        ({}, {"speech_huella": "0" * 64}, 409),
        ({"cargados_mes": 2_500_000}, {}, 409),
        ({"cargados_mes": 2_500_000}, {"confirmar_limite": True}, 201),
        ({"supervisores": ()}, {}, 400),
        ({}, {"speech_huella": "corta"}, 422),
    ],
)
def test_respuestas_de_la_creacion(entorno, extra, estado) -> None:
    client = _cliente(Entorno(**entorno))

    assert _crear(client, **extra).status_code == estado


def test_crear_sin_huella_no_llega_al_caso_de_uso() -> None:
    client = _cliente(Entorno())

    assert client.post("/mowa-mes/campanas", json=CUERPO).status_code == 422


def test_limite_mensual_de_un_mes() -> None:
    client = _cliente(Entorno(cargados_mes=100))

    cuerpo = client.get("/mowa-mes/limite-mensual", params={"mes": "2026-09"}).json()

    assert cuerpo == {
        "mes": "2026-09",
        "limite": 2_500_000,
        "cargados_mes": 100,
        "esta_campana": 0,
        "total": 100,
        "disponible": 2_499_900,
        "excedido": False,
    }
    assert client.get("/mowa-mes/limite-mensual").json()["mes"] == "2026-09"
    assert client.get("/mowa-mes/limite-mensual", params={"mes": "2026-13"}).status_code == 422


# --- Descarga -----------------------------------------------------------


def test_la_descarga_entrega_el_xlsx_y_las_cabeceras_declaradas() -> None:
    client = _cliente(Entorno())
    campana_id = _crear(client).json()["id"]

    respuesta = client.get(f"/mowa-mes/campanas/{campana_id}/archivos/1")

    assert respuesta.status_code == 200
    assert 'filename="mowa_mes_campana_1_1_de_1.xlsx"' in respuesta.headers["content-disposition"]
    enviadas = {n for n in respuesta.headers if n.startswith("x-mowa-mes-")}
    assert enviadas == {nombre.lower() for nombre in CABECERAS_ARCHIVO}
    assert (respuesta.headers["x-mowa-mes-filas"], respuesta.headers["x-mowa-mes-supervision"]) == (
        "6",
        "5",
    )
    hoja = openpyxl.load_workbook(io.BytesIO(respuesta.content))["Hoja1"]
    assert next(hoja.iter_rows(values_only=True)) == ("numero", "mensaje", "dni")
    assert client.get(f"/mowa-mes/campanas/{campana_id}/archivos/2").status_code == 404
    assert client.get("/mowa-mes/campanas/99/archivos/1").status_code == 404


def test_el_contrato_declara_las_cabeceras_de_la_descarga() -> None:
    esquema = contrato.generar()
    respuesta = esquema["paths"]["/mowa-mes/campanas/{campana_id}/archivos/{numero}"]["get"][
        "responses"
    ]["200"]

    assert set(respuesta["headers"]) == {"Content-Disposition", *CABECERAS_ARCHIVO}
    assert list(respuesta["content"]) == [
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ]


# --- Reporte de enviados ------------------------------------------------


def _reporte(filas, mes_id: int) -> bytes:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Hoja1"
    hoja.append(list(COLUMNAS_REPORTE))
    for fila in filas:
        hoja.append(
            [mes_id, fila.numero, fila.mensaje, "14/09/26", fila.dni, "enviado", "Nro. LARGO", "u"]
        )
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def test_importar_el_reporte_y_conciliar() -> None:
    entorno = Entorno()
    client = _cliente(entorno)
    campana_id = _crear(client).json()["id"]
    contenido = _reporte(entorno.campanas.filas_cargadas(campana_id), 990001)

    def importar(**datos):
        return client.post(
            f"/mowa-mes/campanas/{campana_id}/reportes",
            files={"archivo": ("reporte.xlsx", contenido)},
            data=datos,
        )

    importado = importar()
    repetido = importar()
    reemplazado = importar(reemplazar="true")
    conciliacion = client.get(f"/mowa-mes/campanas/{campana_id}/conciliacion").json()

    assert importado.status_code == 201
    assert repetido.status_code == 409
    assert reemplazado.status_code == 201
    assert (conciliacion["total"]["cargados"], conciliacion["total"]["enviados"]) == (6, 6)
    assert conciliacion["supervision"]["por_estado"] == [{"estado": "enviado", "cantidad": 5}]
    assert conciliacion["por_id"] == [{"mes_id": 990001, "filas": 6, "con_correspondencia": 6}]
    assert [r["mes_id"] for r in conciliacion["reportes"]] == [990001]
    assert conciliacion["advertencias"] == []


def test_importar_rechazos() -> None:
    client = _cliente(Entorno())
    campana_id = _crear(client).json()["id"]

    invalido = client.post(
        f"/mowa-mes/campanas/{campana_id}/reportes", files={"archivo": ("x.xlsx", b"no es xlsx")}
    )
    sin_campana = client.post(
        "/mowa-mes/campanas/99/reportes", files={"archivo": ("x.xlsx", b"no es xlsx")}
    )

    assert invalido.status_code == 400
    assert sin_campana.status_code == 404
    assert client.get("/mowa-mes/campanas/99/conciliacion").status_code == 404


# --- Decisiones E-1 y E-2 -----------------------------------------------


def test_la_previsualizacion_advierte_un_documento_no_estandar() -> None:
    carne = producto(2, documento_numero="ZZCE123456", documento_tipo="extranjero")
    client = _cliente(Entorno(productos=[producto(1), carne]))

    cuerpo = client.post("/mowa-mes/campanas/previsualizacion", json=CUERPO).json()

    assert cuerpo["excluidos"] == 0
    assert cuerpo["advertencias_por_codigo"] == [
        {"codigo": "documento_no_estandar", "tipo": "advertencia", "cantidad": 1}
    ]
    productos = [m for m in cuerpo["muestra"] if not m["supervision"]]
    assert [(m["dni"], m["advertencias"]) for m in productos] == [
        ("00000001", []),
        ("ZZCE123456", ["documento_no_estandar"]),
    ]


def test_un_estado_distinto_de_enviado_no_cuenta_en_la_conciliacion() -> None:
    entorno = Entorno()
    client = _cliente(entorno)
    campana_id = _crear(client).json()["id"]
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Hoja1"
    hoja.append(list(COLUMNAS_REPORTE))
    for fila in entorno.campanas.filas_cargadas(campana_id):
        estado = "enviado" if fila.supervision else "fallido"
        hoja.append([990002, fila.numero, fila.mensaje, "14/09/26", fila.dni, estado, "L", "u"])
    buffer = io.BytesIO()
    libro.save(buffer)

    cuerpo = client.post(
        f"/mowa-mes/campanas/{campana_id}/reportes",
        files={"archivo": ("reporte.xlsx", buffer.getvalue())},
    ).json()

    assert cuerpo["productos"] == {
        "cargados": 1,
        "enviados": 0,
        "no_enviados": 1,
        "por_estado": [{"estado": "fallido", "cantidad": 1}],
    }
    assert (cuerpo["total"]["enviados"], cuerpo["total"]["no_enviados"]) == (5, 1)
    assert cuerpo["por_id"] == [{"mes_id": 990002, "filas": 6, "con_correspondencia": 6}]


# --- Revision de la API falsa de F3 (T-MM-B7) ---------------------------


def test_exclusiones_por_codigo_fuera_del_catalogo_o_que_no_es_exclusion() -> None:
    client = _cliente(Entorno())
    campana_id = _crear(client).json()["id"]
    ruta = f"/mowa-mes/campanas/{campana_id}/exclusiones"

    inventado = client.get(ruta, params={"codigo": "codigo_inventado"})
    advertencia = client.get(ruta, params={"codigo": "mensaje_excede_150"})

    assert inventado.status_code == 422
    assert isinstance(inventado.json()["detail"], list)
    assert advertencia.status_code == 200
    assert (advertencia.json()["total"], advertencia.json()["exclusiones"]) == (0, [])


def test_la_conciliacion_sin_reporte_da_todo_cargado_y_nada_enviado() -> None:
    client = _cliente(Entorno())
    campana = _crear(client).json()

    cuerpo = client.get(f"/mowa-mes/campanas/{campana['id']}/conciliacion").json()

    assert cuerpo["reportes"] == []
    for grupo, cargados in (
        ("productos", campana["productos_cargados"]),
        ("supervision", campana["supervision_cargados"]),
        ("total", campana["total_cargados"]),
    ):
        assert cuerpo[grupo] == {
            "cargados": cargados,
            "enviados": 0,
            "no_enviados": cargados,
            "por_estado": [],
        }
    assert (cuerpo["sin_correspondencia"], cuerpo["sin_correspondencia_por_estado"]) == (0, [])
    assert (cuerpo["por_id"], cuerpo["advertencias"]) == ([], [])
