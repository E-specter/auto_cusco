"""Generacion de archivos de carga contra PostgreSQL real, de punta a punta (opcional).

Corre solo con AUTO_CUSCO_DB_TESTS=1 y la base migrada. Recorre la cadena
completa: se ingesta una sabana sintetica por la via normal, se consulta la
cartera con el repositorio de verdad, se aplican las reglas de mapeo y se
escribe el archivo. Los ultimos casos entran por HTTP, sin sustituir ninguna
dependencia, asi que ejercitan todo el camino que usa la interfaz.

Lo que solo aqui se verifica: la paginacion real con LIMIT/OFFSET, el orden que
resuelve PostgreSQL y los tipos que devuelve la base (Decimal y date), que en
las pruebas de nucleo estaban supuestos por un doble.

Fecha del ano 2099 para no cruzarse con datos de trabajo; se borra al terminar.
"""

import io
import os
from datetime import date

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete

from app.adapters.input.lector_calamine import LectorCalamine
from app.adapters.output import exportadores
from app.adapters.persistence.db import get_engine
from app.adapters.persistence.modelos import Carga, CargaAuditoria
from app.adapters.persistence.repositorio_cargas_postgres import RepositorioCargasPostgres
from app.adapters.persistence.repositorio_cartera_postgres import RepositorioCarteraPostgres
from app.core.entities.cartera import Filtro, Operador, Orden
from app.core.entities.mapeo import CampoSalida, DefinicionCarga, TipoSalida
from app.core.services.generacion_cargas import servicio as generacion
from app.core.services.ingesta_sabana.servicio import IngestaSabanaService
from app.core.services.seleccion_cartera.servicio import ConsultaCarteraService
from app.main import app
from tests.test_sabana_cabeceras import CABECERAS_10_09
from tests.test_sabana_normalizador import _fila_sintetica

pytestmark = [
    pytest.mark.postgres,
    pytest.mark.skipif(
        os.environ.get("AUTO_CUSCO_DB_TESTS") != "1",
        reason="Requiere AUTO_CUSCO_DB_TESTS=1 y PostgreSQL migrado",
    ),
]

FECHA = date(2099, 4, 1)
CARGA, AUDITORIA = Carga.__table__, CargaAuditoria.__table__

# Cinco productos con saldos distintos, uno sin telefono valido para ver que un
# campo vacio no rompe la fila. 46278 es el serial de Excel de 2026-09-13.
PRODUCTOS = [
    {
        "Pagare": f"{i:018d}",
        "PAGARE": f"{i:018d}",
        "Región": "CUSCO SUR" if i % 2 else "TACNA",
        "Teléfono": 987654320.0 + i if i != 3 else 12345.0,  # el tercero es invalido
        "Saldo Capital Pendiente": 1000.0 * i + 0.5,
        "Vencimiento Cuota": 46278.0 + i,
    }
    for i in range(1, 6)
]

DEFINICION = DefinicionCarga(
    nombre="SMS preventiva 2099",
    campos=(
        CampoSalida(nombre="anexo", plantilla="1010", tipo=TipoSalida.NUMERO),
        CampoSalida(nombre="numero", plantilla="51[@telefono]"),
        CampoSalida(nombre="pagare", plantilla="[@pagare]"),
        CampoSalida(
            nombre="deuda", plantilla="[@saldo_capital_pendiente]", tipo=TipoSalida.FINANCIERO
        ),
        CampoSalida(
            nombre="vence",
            plantilla="[@fecha_vencimiento_cuota]",
            tipo=TipoSalida.FECHA,
            formato_fecha="%d/%m/%Y",
        ),
    ),
)

DEFINICION_HTTP = {
    "nombre": "SMS preventiva 2099",
    "campos": [
        {"nombre": "anexo", "plantilla": "1010", "tipo": "numero"},
        {"nombre": "numero", "plantilla": "51[@telefono]"},
        {"nombre": "pagare", "plantilla": "[@pagare]"},
        {"nombre": "deuda", "plantilla": "[@saldo_capital_pendiente]", "tipo": "financiero"},
        {
            "nombre": "vence",
            "plantilla": "[@fecha_vencimiento_cuota]",
            "tipo": "fecha",
            "formato_fecha": "%d/%m/%Y",
        },
    ],
}


def _sabana(productos=PRODUCTOS) -> bytes:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "VENCIDA"
    hoja.append(CABECERAS_10_09)
    for cambios in productos:
        hoja.append(_fila_sintetica(**cambios))
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _limpiar(engine) -> None:
    with engine.begin() as cx:
        cx.execute(delete(CARGA).where(CARGA.c.fecha_corte == FECHA))
        cx.execute(delete(AUDITORIA).where(AUDITORIA.c.fecha_corte == FECHA))


@pytest.fixture
def generador():
    """Servicio de generacion apoyado en la cartera real de una sabana ingestada."""
    engine = get_engine()
    _limpiar(engine)
    ingesta = IngestaSabanaService(RepositorioCargasPostgres(engine), LectorCalamine())
    carga = ingesta.registrar_carga(FECHA, "sintetica.xlsx", _sabana()).carga
    assert ingesta.procesar_carga(carga.id).vigente
    consulta = ConsultaCarteraService(RepositorioCarteraPostgres(engine))
    try:
        yield generacion.GeneracionCargasService(consulta), engine
    finally:
        _limpiar(engine)


@pytest.fixture
def cliente(generador):
    """Cliente HTTP sobre la base real: no se sustituye ninguna dependencia."""
    yield TestClient(app)


def _peticion(**extra) -> dict:
    return {
        "fecha_corte": FECHA.isoformat(),
        "definicion": DEFINICION_HTTP,
        "cantidad": 10,
        **extra,
    }


def test_genera_la_carga_desde_la_cartera_real(generador) -> None:
    servicio, _ = generador

    resultado = servicio.generar(FECHA, DEFINICION, cantidad=10)

    assert (resultado.generados, resultado.disponibles) == (5, 5)
    assert resultado.completa
    # Los tipos vienen de PostgreSQL: Decimal se formatea y date se imprime.
    por_pagare = {fila["pagare"]: fila for fila in resultado.tabla.filas}
    assert por_pagare["000000000000000001"]["deuda"] == "1,000.50"
    assert por_pagare["000000000000000001"]["vence"] == "14/09/2026"
    assert por_pagare["000000000000000001"]["numero"] == "51987654321"
    # El tercer producto trae un telefono invalido: la ingesta lo deja vacio y
    # en la carga queda solo el prefijo, sin detener nada.
    assert por_pagare["000000000000000003"]["numero"] == "51"


def test_los_filtros_y_el_orden_los_resuelve_la_base(generador) -> None:
    servicio, _ = generador

    resultado = servicio.generar(
        FECHA,
        DEFINICION,
        [Filtro("region", Operador.IGUAL, ("CUSCO SUR",))],
        Orden("saldo_capital_pendiente", descendente=True),
        cantidad=10,
    )

    assert [fila["pagare"] for fila in resultado.tabla.filas] == [
        "000000000000000005",
        "000000000000000003",
        "000000000000000001",
    ]
    assert resultado.disponibles == 3


def test_la_paginacion_real_no_repite_ni_salta_productos(generador, monkeypatch) -> None:
    servicio, _ = generador
    # Se achica la pagina en vez de ingestar miles de filas: lo que se prueba es
    # el LIMIT/OFFSET real y el desempate por pagare, no el tamano del lote.
    monkeypatch.setattr(generacion, "LIMITE_MAXIMO", 2)
    desplazamientos = _espiar_paginas(monkeypatch)

    resultado = servicio.generar(FECHA, DEFINICION, cantidad=5)

    pagares = [fila["pagare"] for fila in resultado.tabla.filas]
    assert desplazamientos == [0, 2, 4]  # tres consultas, no una
    assert pagares == [f"{i:018d}" for i in range(1, 6)]  # completos, en orden y sin repetir


def _espiar_paginas(monkeypatch) -> list[int]:
    """Anota el desplazamiento de cada consulta, para que la prueba no pase sola."""
    desplazamientos: list[int] = []
    consultar = ConsultaCarteraService.consultar

    def espia(self, *args, **kwargs):
        desplazamientos.append(kwargs["desplazamiento"])
        return consultar(self, *args, **kwargs)

    monkeypatch.setattr(ConsultaCarteraService, "consultar", espia)
    return desplazamientos


def test_pedir_mas_de_los_que_hay_avisa_y_genera_lo_disponible(generador) -> None:
    servicio, _ = generador

    resultado = servicio.generar(FECHA, DEFINICION, cantidad=500)

    assert (resultado.generados, resultado.disponibles) == (5, 5)
    assert not resultado.suficiente  # RF-08


def test_la_generacion_sigue_a_la_version_vigente(generador) -> None:
    servicio, engine = generador
    ingesta = IngestaSabanaService(RepositorioCargasPostgres(engine), LectorCalamine())
    correccion = ingesta.registrar_carga(FECHA, "correccion.xlsx", _sabana(PRODUCTOS[:2])).carga
    ingesta.procesar_carga(correccion.id)

    antes = servicio.generar(FECHA, DEFINICION, cantidad=10).generados
    ingesta.asignar_vigente(correccion.id)
    despues = servicio.generar(FECHA, DEFINICION, cantidad=10).generados

    assert (antes, despues) == (5, 2)


def test_el_xlsx_guarda_cada_valor_con_su_tipo(generador) -> None:
    servicio, _ = generador

    resultado = servicio.generar(FECHA, DEFINICION, cantidad=10)
    archivo = exportadores.exportar(resultado.tabla, "xlsx")

    hoja = openpyxl.load_workbook(io.BytesIO(archivo.contenido)).active
    filas = {fila[2]: fila for fila in hoja.iter_rows(min_row=2, values_only=True)}
    assert archivo.nombre == "SMS_preventiva_2099.xlsx"
    assert [c.value for c in hoja[1]] == ["anexo", "numero", "pagare", "deuda", "vence"]
    # El pagare conserva sus ceros a la izquierda porque se guarda como texto.
    assert "000000000000000001" in filas
    assert filas["000000000000000001"][0] == 1010


# --- De punta a punta por HTTP ----------------------------------------


def test_descargar_el_csv_por_http(cliente) -> None:
    respuesta = cliente.post(
        "/archivos-carga",
        json=_peticion(
            formato="csv",
            filtros=["region:igual:CUSCO SUR"],
            orden="-saldo_capital_pendiente",
            opciones={"delimitador": ";"},
        ),
    )

    lineas = respuesta.content.decode("utf-8").splitlines()
    assert respuesta.status_code == 200
    assert respuesta.headers["content-disposition"].endswith('"SMS_preventiva_2099.csv"')
    assert respuesta.headers["x-carga-generados"] == "3"
    assert respuesta.headers["x-carga-errores"] == "0"
    assert lineas[0] == "anexo;numero;pagare;deuda;vence"
    assert lineas[1] == "1010;51987654325;000000000000000005;5,000.50;18/09/2026"


def test_descargar_el_xlsx_por_http(cliente) -> None:
    respuesta = cliente.post("/archivos-carga", json=_peticion(opciones={"hoja": "Campaña"}))

    hoja = openpyxl.load_workbook(io.BytesIO(respuesta.content)).active
    assert hoja.title == "Campaña"
    assert hoja.max_row == 6  # cabecera mas cinco productos


def test_previsualizar_por_http_antes_de_descargar(cliente) -> None:
    cuerpo = cliente.post("/archivos-carga/previsualizacion", json=_peticion()).json()

    assert cuerpo["cabeceras"] == ["anexo", "numero", "pagare", "deuda", "vence"]
    assert (cuerpo["disponibles"], cuerpo["generados"]) == (5, 5)
    assert cuerpo["completa"] and not cuerpo["suficiente"]  # se pidieron 10 y hay 5
    assert cuerpo["filas"][0]["deuda"] == "1,000.50"


def test_una_fecha_sin_version_vigente_es_404_por_http(cliente) -> None:
    respuesta = cliente.post(
        "/archivos-carga", json=_peticion(fecha_corte=date(2099, 4, 2).isoformat())
    )

    assert respuesta.status_code == 404
