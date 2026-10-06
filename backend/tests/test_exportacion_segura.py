"""Exportacion segura (B3, RF-52, RF-53): un valor exportado no se ejecuta como formula al abrirlo.

Se lee de vuelta cada archivo con `openpyxl` y con el modulo `csv`, y lo que no debe cambiar se
compara contra constantes congeladas en HEAD 3e567d7 (`exportacion_seguridad_congelados.py`).
Todo con datos sinteticos.
"""

import csv
import io
import json
import re
import zipfile
from dataclasses import replace
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import openpyxl.packaging.core
import pytest

from app.adapters.output import exportadores
from app.adapters.output.exportadores import comunes
from app.adapters.output.plataformas.mowa_mes.archivo_carga import escribir_archivo_carga
from app.core.entities.exportacion import OpcionesArchivo, Tabla
from app.core.entities.mowa_mes_campana import FilaCarga
from tests import exportacion_seguridad_congelados as congelados
from tests.exportacion_seguridad_datos import (
    FILAS_MOWA_SIN_IGUAL,
    TABLA_PELIGROSA,
    TABLA_SIN_IGUAL,
    huella_bytes,
    huella_xlsx,
    largo_sin_comprimir_sin_core,
)

SIN_PROTECCION = OpcionesArchivo(proteger_formulas=False)
PUNTO_Y_COMA_CON_BOM = OpcionesArchivo(delimitador=";", codificacion="utf-8-sig")


def _xlsx(tabla: Tabla, opciones: OpcionesArchivo | None = None) -> bytes:
    return exportadores.exportar(tabla, "xlsx", opciones or OpcionesArchivo()).contenido


def _csv(tabla: Tabla, opciones: OpcionesArchivo | None = None) -> list[list[str]]:
    opciones = opciones or OpcionesArchivo()
    archivo = exportadores.exportar(tabla, "csv", opciones)
    texto = archivo.contenido.decode(opciones.codificacion)
    return list(csv.reader(io.StringIO(texto, newline=""), delimiter=opciones.delimitador))


def _hoja(contenido: bytes):
    return openpyxl.load_workbook(io.BytesIO(contenido)).active


def _xml_de_la_hoja(contenido: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(contenido)) as libro:
        return libro.read("xl/worksheets/sheet1.xml").decode("utf-8")


VALORES_DE_TEXTO = [f["valor"] for f in TABLA_PELIGROSA.filas if isinstance(f["valor"], str)]

# Lo que tiene que decir el CSV de la tabla peligrosa, valor por valor, con la proteccion activa.
CSV_ESPERADO = [
    "'=1+1",
    "'=cmd|' /C calc'!A0",
    "'+1+1",
    "'-2+3",
    "'@SUM(A1:A2)",
    "'\t=1+1",
    "'\r=1+1",
    "'-",
    "'+",
    "-1,234.56",  # un importe con formato: no es una formula
    "+51987654321",  # un telefono con prefijo: tampoco
    "-5",
    "3.5",
    "1.234,56",
    "'-.5",  # despues del signo tiene que venir un digito
    "'-5abc",  # todo el texto tiene que ser un numero
    "'= texto con espacio",
    "normal",
    "línea 1\nlínea 2",
    "'=línea 1\nlínea 2",
    "Ñandú =1+1",  # el peligro es al inicio, no en medio
    "",
    "-1234.56",  # Decimal: no es texto, no se toca
    "-5",  # int
    "0",
    "3.5",
    "1",  # True
    "2026-09-18",
    "",  # None
]


# --- XLSX: todo texto es texto literal ---------------------------------------------------------


def test_un_texto_que_empieza_con_igual_ya_no_se_escribe_como_formula() -> None:
    contenido = _xlsx(Tabla("f", ("valor",), ({"valor": "=1+1"},)))

    xml = _xml_de_la_hoja(contenido)

    assert "<f>" not in xml and "<f " not in xml
    assert 't="inlineStr"' in xml and "<t>=1+1</t>" in xml


def test_el_xlsx_devuelve_el_mismo_texto_de_cada_valor_sin_cambiarlo() -> None:
    hoja = _hoja(_xlsx(TABLA_PELIGROSA))

    valores = [fila[0] for fila in hoja.iter_rows(min_row=2)]

    originales = [f["valor"] for f in TABLA_PELIGROSA.filas]
    assert len(valores) >= originales.index(None)  # la hoja no pierde ninguna fila con dato
    for celda, original in zip(valores, originales):  # noqa: B905  (las ultimas filas vacias no se leen)
        if isinstance(original, str) and original != "" and "\r" not in original:
            assert celda.value == original, original
            assert celda.data_type == "s", original


def test_un_retorno_de_carro_en_xlsx_se_lee_como_salto_de_linea_igual_que_antes_de_b3() -> None:
    # openpyxl escribe el \r sin escapar y un lector XML lo normaliza a \n. Ya pasaba en HEAD con
    # cualquier texto: B3 no lo toca. Lo que importa aqui es que siga siendo texto, no formula.
    contenido = _xlsx(Tabla("f", ("v",), ({"v": "\r=1+1"}, {"v": "a\rb"})))

    celdas = [fila[0] for fila in _hoja(contenido).iter_rows(min_row=2)]

    assert [c.value for c in celdas] == ["\n=1+1", "a\nb"]
    assert {c.data_type for c in celdas} == {"s"}
    assert "<f>" not in _xml_de_la_hoja(contenido)


def test_el_xlsx_no_agrega_apostrofo_ni_quote_prefix() -> None:
    contenido = _xlsx(TABLA_PELIGROSA)

    with zipfile.ZipFile(io.BytesIO(contenido)) as libro:
        estilos = libro.read("xl/styles.xml").decode("utf-8")
    xml = _xml_de_la_hoja(contenido)

    assert 'quotePrefix="1"' not in estilos  # el "0" es el valor por defecto de todo estilo
    assert "<t>'=1+1</t>" not in xml
    assert "<t>=1+1</t>" in xml
    assert "<f>" not in xml


def test_el_xlsx_escribe_igual_de_texto_un_igual_en_la_cabecera() -> None:
    contenido = _xlsx(Tabla("f", ("=cabecera", "otra"), ({"=cabecera": "x", "otra": "y"},)))

    assert [c.value for c in _hoja(contenido)[1]] == ["=cabecera", "otra"]
    assert "<f>" not in _xml_de_la_hoja(contenido)


def test_el_xlsx_conserva_los_tipos_nativos() -> None:
    hoja = _hoja(_xlsx(TABLA_PELIGROSA))
    por_fila = {i: fila[0] for i, fila in enumerate(hoja.iter_rows(min_row=2), start=2)}
    originales = [f["valor"] for f in TABLA_PELIGROSA.filas]
    posicion = {
        tipo: originales.index(valor) + 2
        for tipo, valor in {
            "decimal": Decimal("-1234.56"),
            "entero": -5,
            "fecha": date(2026, 9, 18),
        }.items()
    }

    assert por_fila[posicion["decimal"]].data_type == "n"
    assert por_fila[posicion["decimal"]].value == -1234.56
    assert por_fila[posicion["entero"]].value == -5
    assert por_fila[posicion["fecha"]].value == datetime(2026, 9, 18)


def test_el_xlsx_lleva_igual_el_texto_con_saltos_de_linea_dentro_de_la_celda() -> None:
    hoja = _hoja(_xlsx(Tabla("f", ("v",), ({"v": "=línea 1\nlínea 2\n"},))))

    assert hoja["A2"].value == "=línea 1\nlínea 2\n"
    assert hoja["A2"].data_type == "s"


def test_el_xlsx_sin_igual_es_identico_al_de_head_en_todas_sus_partes() -> None:
    contenido = _xlsx(TABLA_SIN_IGUAL)

    assert (huella_xlsx(contenido), largo_sin_comprimir_sin_core(contenido)) == (
        congelados.XLSX_GENERICO_SIN_IGUAL
    )


def _fijar_la_hora_de_openpyxl(monkeypatch, momento: datetime) -> None:
    class Reloj(datetime):
        @classmethod
        def now(cls, tz=None):
            return momento if tz is None else momento.replace(tzinfo=tz)

    monkeypatch.setattr(
        openpyxl.packaging.core, "datetime", SimpleNamespace(datetime=Reloj, timezone=timezone)
    )


def test_la_huella_y_el_largo_congelados_no_dependen_de_la_hora_de_escritura(monkeypatch) -> None:
    # El core.xml lleva la fecha y la hora. Con tres momentos distintos (tambien de dia y de ano),
    # la huella y el largo sin comprimir de las demas partes no se mueven.
    momentos = [datetime(2020, 1, 1, 0, 0, 0), datetime(2031, 12, 31, 23, 59, 59)]
    momentos.append(datetime(2026, 9, 30, 10, 46, 31))
    escritos = []
    for momento in momentos:
        _fijar_la_hora_de_openpyxl(monkeypatch, momento)
        generico = _xlsx(TABLA_SIN_IGUAL)
        mowa = escribir_archivo_carga(FILAS_MOWA_SIN_IGUAL)
        with zipfile.ZipFile(io.BytesIO(generico)) as libro:
            core = libro.read("docProps/core.xml").decode("utf-8")
        escritos.append(
            (
                core,
                huella_xlsx(generico),
                largo_sin_comprimir_sin_core(generico),
                huella_xlsx(mowa),
                largo_sin_comprimir_sin_core(mowa),
            )
        )

    assert len({e[0] for e in escritos}) == 3  # la hora si cambio de verdad
    assert "2020-01-01T00:00:00Z" in escritos[0][0] and "2031-12-31T23:59:59Z" in escritos[1][0]
    assert len({e[1:] for e in escritos}) == 1  # y no cambia nada de lo congelado
    assert escritos[0][1:] == (*congelados.XLSX_GENERICO_SIN_IGUAL, *congelados.XLSX_MOWA_SIN_IGUAL)


def _partes_de(contenido: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(contenido)) as libro:
        return {nombre: libro.read(nombre) for nombre in libro.namelist()}


def test_dos_dias_distintos_solo_cambian_el_core_xml(monkeypatch) -> None:
    # La fecha de escritura va solo en docProps/core.xml: de un dia a otro (misma hora) las demas
    # partes salen byte a byte iguales, y la huella y el largo sin comprimir no se mueven.
    dias = [datetime(2026, 10, 4, 10, 46, 31), datetime(2026, 10, 5, 10, 46, 31)]
    escritos = []
    for dia in dias:
        _fijar_la_hora_de_openpyxl(monkeypatch, dia)
        escritos.append((_xlsx(TABLA_SIN_IGUAL), escribir_archivo_carga(FILAS_MOWA_SIN_IGUAL)))

    for del_primero, del_segundo in zip(*escritos, strict=True):  # el generico y el de MOWA MES
        uno, dos = _partes_de(del_primero), _partes_de(del_segundo)
        assert uno.keys() == dos.keys()
        assert {nombre for nombre in uno if uno[nombre] != dos[nombre]} == {"docProps/core.xml"}
        assert "2026-10-04T10:46:31Z" in uno["docProps/core.xml"].decode("utf-8")
        assert "2026-10-05T10:46:31Z" in dos["docProps/core.xml"].decode("utf-8")
        assert huella_xlsx(del_primero) == huella_xlsx(del_segundo)
        assert largo_sin_comprimir_sin_core(del_primero) == largo_sin_comprimir_sin_core(
            del_segundo
        )


def test_la_proteccion_del_xlsx_no_depende_de_la_opcion() -> None:
    # En XLSX no hay opcion: proteger_formulas=False solo afecta al CSV.
    tabla = Tabla("f", ("v",), ({"v": "=1+1"},))

    assert "<f>" not in _xml_de_la_hoja(_xlsx(tabla, SIN_PROTECCION))


# --- CSV: apostrofo delante de lo que parece una formula ---------------------------------------


def test_el_csv_pone_apostrofo_a_lo_que_empieza_como_formula_y_no_a_lo_demas() -> None:
    filas = _csv(TABLA_PELIGROSA)

    assert filas[0] == ["valor"]
    assert [fila[0] for fila in filas[1:]] == CSV_ESPERADO


def test_el_csv_no_protege_un_importe_ni_un_entero_negativos_que_no_son_texto() -> None:
    tabla = Tabla("f", ("d", "i", "b"), ({"d": Decimal("-1234.56"), "i": -5, "b": True},))

    assert _csv(tabla)[1] == ["-1234.56", "-5", "1"]


def test_el_csv_no_protege_un_decimal_aunque_su_texto_no_sea_un_numero_plano() -> None:
    # str(Decimal("-0.0000001")) es "-1E-7": empieza con "-" y no es solo digitos, pero un Decimal
    # no es un texto, no es una formula y el apostrofo lo rompe.
    valores = [Decimal("-0.0000001"), Decimal("-1E+3"), Decimal("-Infinity"), Decimal("-0")]
    tabla = Tabla("f", ("v",), tuple({"v": v} for v in valores))

    assert [fila[0] for fila in _csv(tabla)[1:]] == ["-1E-7", "-1E+3", "-Infinity", "-0"]


def test_el_csv_protege_un_texto_solo_si_no_es_un_numero_con_signo() -> None:
    tabla = Tabla(
        "f", ("v",), tuple({"v": v} for v in ["-2+3", "-5", "-5.5", "-5,5", "+5-5", "-5\n"])
    )

    assert [f[0] for f in _csv(tabla)[1:]] == ["'-2+3", "-5", "-5.5", "-5,5", "'+5-5", "'-5\n"]


def test_el_csv_protege_los_digitos_que_no_son_ascii_como_no_numericos() -> None:
    tabla = Tabla("f", ("v",), ({"v": "-٣"}, {"v": "+²"}))

    assert [f[0] for f in _csv(tabla)[1:]] == ["'-٣", "'+²"]


def test_el_csv_lleva_bien_tildes_saltos_y_delimitadores_de_un_texto_protegido() -> None:
    tabla = Tabla("f", ("v",), ({"v": '=Ñandú;"x"\nsalto'},))

    assert _csv(tabla, PUNTO_Y_COMA_CON_BOM)[1] == ['\'=Ñandú;"x"\nsalto']


def test_con_la_proteccion_apagada_el_csv_es_identico_al_de_head() -> None:
    for tabla, esperado, esperado_con_bom in (
        (TABLA_PELIGROSA, congelados.CSV_PELIGROSA, congelados.CSV_PELIGROSA_PUNTO_Y_COMA_CON_BOM),
        (TABLA_SIN_IGUAL, congelados.CSV_SIN_IGUAL, congelados.CSV_SIN_IGUAL_PUNTO_Y_COMA_CON_BOM),
    ):
        normal = exportadores.exportar(tabla, "csv", SIN_PROTECCION).contenido
        con_bom = exportadores.exportar(
            tabla, "csv", replace(PUNTO_Y_COMA_CON_BOM, proteger_formulas=False)
        ).contenido

        assert huella_bytes(normal) == esperado
        assert huella_bytes(con_bom) == esperado_con_bom


def test_por_defecto_la_proteccion_esta_activa_y_solo_cambia_lo_peligroso() -> None:
    protegido = exportadores.exportar(TABLA_SIN_IGUAL, "csv").contenido.decode("utf-8")
    antes = exportadores.exportar(TABLA_SIN_IGUAL, "csv", SIN_PROTECCION).contenido.decode("utf-8")

    assert OpcionesArchivo().proteger_formulas is True
    assert huella_bytes(antes.encode("utf-8")) == congelados.CSV_SIN_IGUAL
    diferencias = [
        (a, p)
        for a, p in zip(
            csv.reader(io.StringIO(antes, newline="")),
            csv.reader(io.StringIO(protegido, newline="")),
            strict=True,
        )
        if a != p
    ]
    # Solo cambian las filas con "@usuario", "-2+3", tabulacion y retorno al inicio de un texto.
    assert len(diferencias) == 4
    assert {p[i] for a, p in diferencias for i in range(len(p)) if a[i] != p[i]} == {
        "'@usuario",
        "'-2+3",
        "'\tcon tabulacion",
        "'\rcon retorno",
    }


def test_las_cabeceras_del_csv_se_protegen_con_el_mismo_criterio() -> None:
    # Una definicion compartida (RF-51) puede traer un nombre de columna "=...": es un valor
    # exportado como cualquier otro y RF-53 no distingue.
    cabeceras = (
        "=a",
        "-b",
        "+c",
        "@d",
        chr(9) + "e",
        chr(13) + "f",
        "-5",
        "+51987654321",
        "normal",
    )
    tabla = Tabla("f", cabeceras, ({c: "x" for c in cabeceras},))

    filas = _csv(tabla)

    assert filas[0] == [
        "'=a",
        "'-b",
        "'+c",
        "'@d",
        "'" + chr(9) + "e",
        "'" + chr(13) + "f",
        "-5",
        "+51987654321",
        "normal",
    ]
    assert filas[1] == ["x"] * len(cabeceras)  # los valores "x" no se tocan


def test_con_la_proteccion_apagada_las_cabeceras_del_csv_no_cambian() -> None:
    tabla = Tabla("f", ("=a", "-b"), ({"=a": "x", "-b": "y"},))

    assert _csv(tabla, SIN_PROTECCION)[0] == ["=a", "-b"]


def test_la_cabecera_del_csv_es_texto_y_un_numero_con_signo_no_se_protege() -> None:
    tabla = Tabla("f", ("-1,234.56", "-2+3"), ({"-1,234.56": "x", "-2+3": "y"},))

    assert _csv(tabla)[0] == ["-1,234.56", "'-2+3"]


def test_el_csv_sin_cabeceras_protege_igual_las_filas() -> None:
    tabla = Tabla("f", ("v",), ({"v": "=1+1"},))

    assert _csv(tabla, OpcionesArchivo(con_cabeceras=False)) == [["'=1+1"]]


def test_el_csv_no_usa_expresiones_regulares_por_celda(monkeypatch) -> None:
    def _sin_regex(*args, **kwargs):
        raise AssertionError("se compilo una expresion regular al exportar")

    monkeypatch.setattr(re, "_compile", _sin_regex)

    filas = _csv(TABLA_PELIGROSA)

    assert len(filas) == len(TABLA_PELIGROSA.filas) + 1


# --- JSON: sin cambios --------------------------------------------------------------------------


@pytest.mark.parametrize("opciones", [OpcionesArchivo(), SIN_PROTECCION])
def test_el_json_es_identico_al_de_head_con_o_sin_proteccion(opciones) -> None:
    peligrosa = exportadores.exportar(TABLA_PELIGROSA, "json", opciones).contenido
    sin_igual = exportadores.exportar(TABLA_SIN_IGUAL, "json", opciones).contenido

    assert huella_bytes(peligrosa) == congelados.JSON_PELIGROSA
    assert huella_bytes(sin_igual) == congelados.JSON_SIN_IGUAL
    assert json.loads(peligrosa)[0]["valor"] == "=1+1"  # el JSON no es una hoja: sin apostrofo


# --- La opcion en la API ------------------------------------------------------------------------


def test_la_opcion_la_fija_la_definicion_y_no_las_opciones_de_archivos_carga() -> None:
    """D4: `proteger_formulas` viaja en la definicion (la que se arma, la que se guarda y la que
    se devuelve) y pasa a las opciones del archivo al generar con una guardada; `/archivos-carga`
    no la expone como opcion propia (la definicion en linea sigue protegiendo por defecto)."""
    contrato = Path(__file__).resolve().parents[2] / "contratos" / "openapi.json"

    assert contrato.exists(), f"No se encontro {contrato}"
    esquemas = json.loads(contrato.read_text(encoding="utf-8"))["components"]["schemas"]
    con_la_opcion = sorted(
        n for n, e in esquemas.items() if "proteger_formulas" in e.get("properties", {})
    )

    assert con_la_opcion == ["ContenidoGuardadoRespuesta", "DefinicionColumnasEntrada"]


def test_el_valor_por_defecto_es_proteger() -> None:
    assert OpcionesArchivo().proteger_formulas is True
    assert OpcionesArchivo(proteger_formulas=False).proteger_formulas is False


def test_los_helpers_de_comunes_no_cambian_los_valores_que_no_son_texto() -> None:
    assert comunes.a_texto(Decimal("-1234.56")) == "-1234.56"
    assert comunes.a_texto(True) == "1"
    assert comunes.a_texto(None) == ""


# --- MOWA MES: el archivo de carga -------------------------------------------------------------


def _filas_mowa_con_igual() -> tuple[FilaCarga, ...]:
    return (
        FilaCarga("900000001", "ZZPRUEBA supervision", "00000001", supervision=True),
        FilaCarga("900000101", "=ZZPRUEB vencio el 06/09/2026", "01234567"),
        FilaCarga("900000102", "-ZZPRUEB vencio el 06/09/2026", "20123456789"),
        FilaCarga("900000103", "=1+1", "00000009"),
        FilaCarga("900000104", "+ZZPRUEB acércate\na pagar", "1"),
        FilaCarga("900000105", "@ZZPRUEB, Ñandú", "12345678"),
    )


def test_el_archivo_de_mowa_mes_devuelve_mensaje_y_dni_identicos_a_las_celdas() -> None:
    filas = _filas_mowa_con_igual()

    hoja = _hoja(escribir_archivo_carga(filas))
    leidas = list(hoja.iter_rows(min_row=2))

    assert [c.value for c in hoja[1]] == ["numero", "mensaje", "dni"]
    assert len(leidas) == len(filas)
    for (numero, mensaje, dni), fila in zip(leidas, filas, strict=True):
        esperado = fila.celdas()
        assert numero.value == esperado["numero"] and numero.data_type == "n"
        assert mensaje.value == esperado["mensaje"] and len(mensaje.value) == len(fila.mensaje)
        assert mensaje.data_type == "s"
        assert dni.value == esperado["dni"] and dni.data_type == "s"


def test_el_archivo_de_mowa_mes_no_tiene_formulas_aunque_un_mensaje_empiece_con_igual() -> None:
    xml = _xml_de_la_hoja(escribir_archivo_carga(_filas_mowa_con_igual()))

    assert "<f>" not in xml and "<f " not in xml
    assert "<t>=ZZPRUEB vencio el 06/09/2026</t>" in xml
    assert "<t>=1+1</t>" in xml


def test_el_archivo_de_mowa_mes_sin_igual_es_identico_al_de_head_y_mide_lo_mismo() -> None:
    contenido = escribir_archivo_carga(FILAS_MOWA_SIN_IGUAL)

    assert (
        huella_xlsx(contenido),
        largo_sin_comprimir_sin_core(contenido),
    ) == congelados.XLSX_MOWA_SIN_IGUAL


def test_la_hoja_y_las_cabeceras_de_mowa_mes_no_cambian() -> None:
    libro = openpyxl.load_workbook(io.BytesIO(escribir_archivo_carga(FILAS_MOWA_SIN_IGUAL)))

    assert libro.sheetnames == ["Hoja1"]
