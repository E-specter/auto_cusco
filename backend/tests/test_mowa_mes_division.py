"""Pruebas de la division de la carga (RF-MM-11) y del archivo .xlsx de MES (RF-MM-10)."""

import io

import openpyxl
import pytest

from app.adapters.output.plataformas.mowa_mes.archivo_carga import escribir_archivo_carga
from app.core.entities.mowa_mes_campana import ArchivoDemasiadoGrande, FilaCarga
from app.core.services.plataformas.mowa_mes import division


def _filas(cantidad: int, supervision: int = 0) -> list[FilaCarga]:
    return [
        FilaCarga(
            numero=f"9000{i:05d}",
            mensaje=f"ZZPRUEBA mensaje sintetico {i}",
            dni=f"{i:08d}",
            supervision=i < supervision,
        )
        for i in range(cantidad)
    ]


def _escritor_de_documentos(filas) -> bytes:
    """Cada fila ocupa 9 bytes (su dni y una coma): el tamano es predecible."""
    return ",".join(fila.dni for fila in filas).encode()


def test_tramos_por_filas() -> None:
    assert division.tramos_por_filas(0, 3) == []
    assert division.tramos_por_filas(7, 3) == [(0, 3), (3, 6), (6, 7)]
    assert division.tramos_por_filas(6, 3) == [(0, 3), (3, 6)]
    with pytest.raises(ValueError):
        division.tramos_por_filas(5, 0)


def test_primero_divide_por_filas_contando_la_supervision_en_el_primer_archivo() -> None:
    filas = _filas(7, supervision=2)

    archivos = division.dividir(filas, 3, 10_000, _escritor_de_documentos)

    assert [(a.numero, a.filas, a.supervision) for a in archivos] == [
        (1, 3, 2),
        (2, 3, 0),
        (3, 1, 0),
    ]


def test_si_un_tramo_pasa_de_los_bytes_se_parte_y_conserva_el_orden() -> None:
    filas = _filas(10)

    archivos = division.dividir(filas, 50, 30, _escritor_de_documentos)

    assert all(a.bytes <= 30 for a in archivos)
    assert len(archivos) > 1
    assert [a.numero for a in archivos] == list(range(1, len(archivos) + 1))
    recompuesto = ",".join(a.contenido.decode() for a in archivos).split(",")
    assert recompuesto == [fila.dni for fila in filas]
    assert sum(a.filas for a in archivos) == 10


def test_sin_limite_alcanzado_sale_un_solo_archivo() -> None:
    archivos = division.dividir(_filas(5, supervision=5), 50, 10_000, _escritor_de_documentos)

    assert [(a.filas, a.supervision) for a in archivos] == [(5, 5)]


def test_una_fila_mas_grande_que_el_maximo_no_se_puede_dividir() -> None:
    with pytest.raises(ArchivoDemasiadoGrande):
        division.dividir(_filas(2), 50, 5, _escritor_de_documentos)


def test_archivos_previstos_por_filas() -> None:
    previstos = division.previstos_por_filas(_filas(7, supervision=2), 3)

    assert [(p.numero, p.filas, p.supervision) for p in previstos] == [
        (1, 3, 2),
        (2, 3, 0),
        (3, 1, 0),
    ]


# --- Archivo .xlsx de MES -------------------------------------------------


def test_el_archivo_de_carga_tiene_el_formato_de_mes() -> None:
    filas = [
        FilaCarga("900000001", "ZZPRUEBA supervision", "00000001", supervision=True),
        FilaCarga("900000101", "ZZPRUEBA vencio el 06/09/2026", "01234567"),
        FilaCarga("900000102", "ZZPRUEBA acércate a pagar", "20123456789"),
    ]

    libro = openpyxl.load_workbook(io.BytesIO(escribir_archivo_carga(filas)))

    assert libro.sheetnames == ["Hoja1"]
    hoja = libro["Hoja1"]
    valores = list(hoja.iter_rows(values_only=True))
    assert valores[0] == ("numero", "mensaje", "dni")
    assert valores[1] == (900000001, "ZZPRUEBA supervision", "00000001")
    assert isinstance(valores[2][0], int)
    assert valores[2][2] == "01234567"  # texto: conserva el cero
    assert valores[3] == (900000102, "ZZPRUEBA acércate a pagar", "20123456789")
