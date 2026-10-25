"""Procesar una carga con el mapeo de encabezados confirmado (T-MAP-B3a, RF-57, RF-60 a RF-62).

Sin base de datos: el repositorio en memoria de la ingesta y el lector real. Archivos sinteticos.
"""

import io
import json
from datetime import UTC, date, datetime

import openpyxl
import pytest

from app.adapters.input.lector_tabla import LectorTabla
from app.core.entities.carga import EstadoCarga
from app.core.entities.lectura_tabla import OpcionesLectura
from app.core.entities.mapeador import Asignacion, EstadoEncabezado, OrigenAsignacion
from app.core.entities.sabana import Severidad
from app.core.services.borradores_carga.servicio import MapeadorService
from app.core.services.ingesta_sabana.mapeable import CATALOGO_SABANA_MAPEABLE
from app.core.services.ingesta_sabana.servicio import IngestaSabanaService
from tests.test_ingesta_servicio import RepositorioEnMemoria
from tests.test_mapeador_ciclo import SABANA, aceptar_todo, csv_de
from tests.test_sabana_cabeceras import CABECERAS_10_09
from tests.test_sabana_normalizador import _fila_sintetica

O = OrigenAsignacion  # noqa: E741
E = EstadoEncabezado
FECHA = date(2026, 9, 10)
OTRA_FECHA = date(2026, 9, 11)
INICIO = datetime(2026, 10, 26, 9, 0, tzinfo=UTC)


class Espia(LectorTabla):
    """El lector real que anota con que opciones se le pidio leer."""

    def __init__(self) -> None:
        super().__init__()
        self.pedidos: list[tuple[str, OpcionesLectura]] = []

    def leer(self, contenido, nombre_archivo, opciones=OpcionesLectura()):  # noqa: B008
        self.pedidos.append((nombre_archivo, opciones))
        return super().leer(contenido, nombre_archivo, opciones)


class Entorno:
    def __init__(self) -> None:
        self.repo = RepositorioEnMemoria()
        self.lector = Espia()
        self.mapeador = MapeadorService(
            self.repo, self.lector, [CATALOGO_SABANA_MAPEABLE], lambda: INICIO
        )
        self.ingesta = IngestaSabanaService(self.repo, self.lector, lambda: INICIO)

    def confirmar(self, nombre, contenido, fecha=FECHA, cambios=None, **opciones):
        analizado = self.mapeador.crear_borrador(
            SABANA, nombre, contenido, OpcionesLectura(**opciones)
        )
        asignaciones = aceptar_todo(analizado)
        for indice, campo in (cambios or {}).items():
            asignaciones[indice] = Asignacion(indice, campo, O.MANUAL)
        return self.ingesta.registrar_desde_borrador(
            analizado.borrador.id, fecha, asignaciones
        ).carga

    def fallos(self, carga_id: int) -> str | None:
        return self.repo.cargas[carga_id].get("motivo_fallo")


@pytest.fixture
def entorno() -> Entorno:
    return Entorno()


def xlsx(filas: list[list], hoja: str = "VENCIDA") -> bytes:
    libro = openpyxl.Workbook()
    libro.active.title = hoja
    for fila in filas:
        libro.active.append(fila)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


CSV = csv_de(
    [
        ["Pagare", "Titular", "Teléfono", "Región", "Observación libre"],
        [
            "000000000000000001",
            "APELLIDO/NOMBRE FICTICIO",
            "987654321",
            "REGIÓN PRUEBA",
            "nota uno",
        ],
        ["000000000000000002", "OTRO/FICTICIO", "912345678", "REGIÓN PRUEBA", "nota dos"],
    ]
)


# --- Un CSV con su separador y su codificacion (M-2) --------------------------------------------


def test_un_csv_confirmado_se_procesa_con_su_separador_y_su_codificacion(entorno) -> None:
    carga = entorno.confirmar("datos.csv", CSV)

    resultado = entorno.ingesta.procesar_carga(carga.id)

    assert resultado.estado is EstadoCarga.TERMINADA
    nombre, opciones = entorno.lector.pedidos[-1]
    assert nombre == "datos.csv"
    assert opciones == OpcionesLectura(hoja=None, separador=";", codificacion="windows-1252")
    filas = entorno.repo.filas[carga.id]
    assert [f.valores["pagare"] for f in filas] == ["000000000000000001", "000000000000000002"]
    assert filas[0].valores["region"] == "REGIÓN PRUEBA"  # la tilde sobrevive a Windows-1252
    assert [f.numero_fila for f in filas] == [2, 3]


def test_lo_elegido_a_mano_se_guarda_y_se_usa_al_procesar(entorno) -> None:
    contenido = csv_de([["Pagare", "Titular"], ["000000000000000001", "UNO"]], separador=",")
    carga = entorno.confirmar("datos.csv", contenido, separador=",", codificacion="latin-1")

    entorno.ingesta.procesar_carga(carga.id)

    assert entorno.repo.cargas[carga.id]["opciones_lectura"] == {
        "formato": "csv",
        "hoja": None,
        "separador": ",",
        "codificacion": "latin-1",
    }
    assert entorno.lector.pedidos[-1][1] == OpcionesLectura(
        hoja=None, separador=",", codificacion="latin-1"
    )


def test_una_hoja_elegida_se_procesa_con_esa_hoja(entorno) -> None:
    libro = openpyxl.Workbook()
    libro.active.title = "Resumen"
    libro.active.append(["nada"])
    datos = libro.create_sheet("Septiembre")
    datos.append(["Pagare", "Titular"])
    datos.append(["000000000000000001", "UNO"])
    buffer = io.BytesIO()
    libro.save(buffer)

    carga = entorno.confirmar("sabana.xlsx", buffer.getvalue(), hoja="Septiembre")

    assert entorno.ingesta.procesar_carga(carga.id).estado is EstadoCarga.TERMINADA
    assert entorno.repo.cargas[carga.id]["hoja"] == "Septiembre"
    assert entorno.repo.cargas[carga.id]["opciones_lectura"]["hoja"] == "Septiembre"


# --- POST /cargas con archivo: igual que hoy, para cualquier formato ----------------------------


def test_registrar_un_archivo_guarda_la_hoja_como_opcion_de_lectura(entorno) -> None:
    carga = entorno.ingesta.registrar_carga(FECHA, "sabana.xlsx", b"x").carga

    assert entorno.repo.cargas[carga.id]["opciones_lectura"] == {"hoja": "VENCIDA"}
    assert entorno.repo.cargas[carga.id].get("asignaciones") is None


def test_un_csv_subido_sin_pasos_se_detecta_y_mapea_por_alias(entorno) -> None:
    carga = entorno.ingesta.registrar_carga(FECHA, "datos.csv", CSV).carga

    resultado = entorno.ingesta.procesar_carga(carga.id)

    assert resultado.estado is EstadoCarga.TERMINADA
    assert entorno.repo.filas[carga.id][0].extras == {"Observación libre": "nota uno"}
    mapeo = entorno.repo.cargas[carga.id]["formato"].mapeo
    assert set(mapeo) == {"columnas", "desconocidas", "faltantes"}  # sin origenes: no hubo pasos
    assert mapeo["desconocidas"] == {"5": "Observación libre"}


def test_un_json_subido_avisa_del_valor_no_escalar_en_su_fila(entorno) -> None:
    objetos = [
        {"Pagare": "000000000000000001", "Titular": "UNO", "Detalle": {"a": 1, "b": [2]}},
        {"Pagare": "000000000000000002", "Titular": "DOS", "Detalle": "texto"},
    ]
    carga = entorno.ingesta.registrar_carga(FECHA, "datos.json", json.dumps(objetos).encode()).carga

    resultado = entorno.ingesta.procesar_carga(carga.id)

    assert resultado.estado is EstadoCarga.TERMINADA
    avisos = [i for i in entorno.repo.incidencias[carga.id] if i.codigo == "valor_no_escalar"]
    assert [(i.fila, i.columna, i.severidad) for i in avisos] == [
        (2, "Detalle", Severidad.ADVERTENCIA)
    ]
    assert entorno.repo.filas[carga.id][0].extras == {"Detalle": '{"a": 1, "b": [2]}'}
    assert entorno.repo.filas[carga.id][1].extras == {"Detalle": "texto"}


# --- «No cargar» va al registro crudo (RF-60) ---------------------------------------------------


def test_las_columnas_en_no_cargar_van_a_los_extras_de_cada_fila(entorno) -> None:
    carga = entorno.confirmar("datos.csv", CSV, cambios={3: None})  # «Región» no se carga

    entorno.ingesta.procesar_carga(carga.id)

    fila = entorno.repo.filas[carga.id][0]
    assert fila.valores.get("region") is None
    assert fila.extras == {"Región": "REGIÓN PRUEBA", "Observación libre": "nota uno"}
    mapeo = entorno.repo.cargas[carga.id]["formato"].mapeo
    assert mapeo["ignoradas"] == {"4": "Región", "5": "Observación libre"}
    assert "region" in mapeo["faltantes"]


def test_un_campo_obligatorio_nunca_queda_sin_columna_al_procesar(entorno) -> None:
    carga = entorno.confirmar("datos.csv", CSV)

    entorno.ingesta.procesar_carga(carga.id)

    assert entorno.repo.cargas[carga.id]["estado"] is EstadoCarga.TERMINADA


# --- Los origenes y las columnas ignoradas en formato.mapeo (RF-62) -----------------------------


def test_formato_mapeo_suma_origenes_e_ignoradas_sin_cambiar_lo_demas(entorno) -> None:
    primero = entorno.confirmar("datos.csv", CSV)
    entorno.ingesta.procesar_carga(primero.id)
    segundo = entorno.confirmar("datos.csv", CSV, fecha=OTRA_FECHA)  # ya con lo recordado
    entorno.ingesta.procesar_carga(segundo.id)

    primero_mapeo = entorno.repo.cargas[primero.id]["formato"].mapeo
    segundo_mapeo = entorno.repo.cargas[segundo.id]["formato"].mapeo

    assert primero_mapeo["origenes"] == {
        "pagare": "coincide",
        "titular": "coincide",
        "telefono": "coincide",
        "region": "coincide",
    }
    assert segundo_mapeo["origenes"] == dict.fromkeys(
        ["pagare", "titular", "telefono", "region"], "recordado"
    )
    assert primero_mapeo["columnas"] == {"pagare": 1, "titular": 2, "telefono": 3, "region": 4}
    assert primero_mapeo["ignoradas"] == {"5": "Observación libre"}


# --- Paridad con el camino de siempre -----------------------------------------------------------


def _libro_de_sabana(extra: bool) -> bytes:
    cabeceras = [*CABECERAS_10_09, *(["Columna inventada"] if extra else [])]
    fila = [*_fila_sintetica(), *(["dato suelto"] if extra else [])]
    segunda = [*_fila_sintetica(Pagare="000000000000000002"), *(["otro suelto"] if extra else [])]
    return xlsx([cabeceras, fila, segunda])


@pytest.mark.parametrize("con_columna_desconocida", [False, True])
def test_un_archivo_que_carga_por_alias_da_la_misma_version_por_los_dos_caminos(
    entorno, con_columna_desconocida
) -> None:
    contenido = _libro_de_sabana(con_columna_desconocida)
    por_alias = entorno.ingesta.registrar_carga(FECHA, "sabana.xlsx", contenido).carga
    desde_borrador = entorno.confirmar("sabana.xlsx", contenido, fecha=OTRA_FECHA)

    resultado_a = entorno.ingesta.procesar_carga(por_alias.id)
    resultado_b = entorno.ingesta.procesar_carga(desde_borrador.id)

    assert resultado_a.estado is resultado_b.estado is EstadoCarga.TERMINADA
    assert resultado_a.resumen == resultado_b.resumen
    assert entorno.repo.filas[por_alias.id] == entorno.repo.filas[desde_borrador.id]
    assert entorno.repo.incidencias[por_alias.id] == entorno.repo.incidencias[desde_borrador.id]
    formato_a = entorno.repo.cargas[por_alias.id]["formato"]
    formato_b = entorno.repo.cargas[desde_borrador.id]["formato"]
    assert formato_a.fila_cabecera == formato_b.fila_cabecera
    assert formato_a.cabeceras_originales == formato_b.cabeceras_originales
    assert formato_a.huella_formato == formato_b.huella_formato
    for clave in ("columnas", "desconocidas", "faltantes"):  # las claves de hoy, intactas
        assert formato_a.mapeo[clave] == formato_b.mapeo[clave]
    assert set(formato_a.mapeo) == {"columnas", "desconocidas", "faltantes"}
    assert set(formato_b.mapeo) == {
        "columnas",
        "desconocidas",
        "faltantes",
        "origenes",
        "ignoradas",
    }
    assert formato_b.mapeo["ignoradas"] == formato_b.mapeo["desconocidas"]


def test_la_paridad_incluye_la_columna_duplicada_del_pagare(entorno) -> None:
    cabeceras = ["Pagare", "Titular", "PAGARE", "Pagare"]  # tres: la tercera repite
    contenido = xlsx([cabeceras, ["000000000000000001", "UNO", "000000000000000001", "x"]])
    por_alias = entorno.ingesta.registrar_carga(FECHA, "a.xlsx", contenido).carga
    desde_borrador = entorno.confirmar("a.xlsx", contenido, fecha=OTRA_FECHA)

    entorno.ingesta.procesar_carga(por_alias.id)
    entorno.ingesta.procesar_carga(desde_borrador.id)

    assert entorno.repo.incidencias[por_alias.id] == entorno.repo.incidencias[desde_borrador.id]
    assert any(i.codigo == "cabecera_duplicada" for i in entorno.repo.incidencias[por_alias.id])


# --- El mapeo guardado tiene que seguir siendo el del archivo -----------------------------------


def _cambiar_archivo(entorno: Entorno, carga_id: int, contenido: bytes) -> None:
    entorno.repo.archivos[carga_id] = contenido


def test_un_encabezado_cambiado_en_la_misma_posicion_falla_diciendo_cual(entorno) -> None:
    carga = entorno.confirmar("datos.csv", CSV)
    _cambiar_archivo(
        entorno,
        carga.id,
        CSV.replace("Titular".encode("cp1252"), "Nombre completo".encode("cp1252")),
    )

    resultado = entorno.ingesta.procesar_carga(carga.id)

    assert resultado.estado is EstadoCarga.FALLIDA
    assert "columna 2" in resultado.motivo_fallo
    assert "'titular'" in resultado.motivo_fallo and "'nombre completo'" in resultado.motivo_fallo
    assert entorno.repo.filas.get(carga.id) is None  # no cargo nada en el campo equivocado


def test_un_indice_fuera_de_rango_falla_con_un_motivo_claro(entorno) -> None:
    carga = entorno.confirmar("datos.csv", CSV)
    _cambiar_archivo(
        entorno, carga.id, csv_de([["Pagare", "Titular"], ["000000000000000001", "UNO"]])
    )

    resultado = entorno.ingesta.procesar_carga(carga.id)

    assert resultado.estado is EstadoCarga.FALLIDA
    assert "columna 3" in resultado.motivo_fallo and "2 columnas" in resultado.motivo_fallo


def test_una_columna_nueva_sin_asignacion_falla(entorno) -> None:
    carga = entorno.confirmar("datos.csv", CSV)
    _cambiar_archivo(
        entorno,
        carga.id,
        csv_de(
            [
                ["Pagare", "Titular", "Teléfono", "Región", "Observación libre", "Columna nueva"],
                ["000000000000000001", "UNO", "987654321", "R", "n", "x"],
            ]
        ),
    )

    resultado = entorno.ingesta.procesar_carga(carga.id)

    assert resultado.estado is EstadoCarga.FALLIDA
    assert "columna 6" in resultado.motivo_fallo and "Columna nueva" in resultado.motivo_fallo


def test_el_caso_normal_pasa_la_comprobacion(entorno) -> None:
    carga = entorno.confirmar("datos.csv", CSV)

    assert entorno.ingesta.procesar_carga(carga.id).estado is EstadoCarga.TERMINADA


def test_el_mismo_numero_de_columnas_no_basta(entorno) -> None:
    # Dos columnas intercambiadas: la cantidad no cambia, pero la columna 2 ya no es «Titular».
    carga = entorno.confirmar("datos.csv", CSV)
    intercambiado = csv_de(
        [
            ["Pagare", "Teléfono", "Titular", "Región", "Observación libre"],
            ["000000000000000001", "987654321", "UNO", "REGIÓN PRUEBA", "nota"],
        ]
    )
    _cambiar_archivo(entorno, carga.id, intercambiado)

    assert entorno.ingesta.procesar_carga(carga.id).estado is EstadoCarga.FALLIDA


def test_un_campo_que_ya_no_esta_en_el_catalogo_falla(entorno) -> None:
    carga = entorno.confirmar("datos.csv", CSV)
    entorno.repo.cargas[carga.id]["asignaciones"][1]["campo"] = "campo_retirado"

    resultado = entorno.ingesta.procesar_carga(carga.id)

    assert resultado.estado is EstadoCarga.FALLIDA and "campo_retirado" in resultado.motivo_fallo


def test_una_carga_fallida_por_el_mapeo_conserva_el_formato_leido(entorno) -> None:
    carga = entorno.confirmar("datos.csv", CSV)
    _cambiar_archivo(entorno, carga.id, csv_de([["Pagare"], ["000000000000000001"]]))

    entorno.ingesta.procesar_carga(carga.id)

    formato = entorno.repo.cargas[carga.id]["formato"]
    assert formato.cabeceras_originales == ["Pagare"] and formato.fila_cabecera == 1


# --- Reprocesar usa las mismas asignaciones (RF-33) ---------------------------------------------


def test_reprocesar_una_carga_interrumpida_usa_las_mismas_asignaciones(entorno) -> None:
    carga = entorno.confirmar("datos.csv", CSV, cambios={3: None})
    asignaciones = list(entorno.repo.cargas[carga.id]["asignaciones"])
    entorno.repo.cargas[carga.id]["estado"] = EstadoCarga.PROCESANDO  # se cayo a mitad

    assert entorno.ingesta.recuperar_interrumpidas() == [carga.id]
    resultado = entorno.ingesta.procesar_carga(carga.id)

    assert resultado.estado is EstadoCarga.TERMINADA
    assert entorno.repo.cargas[carga.id]["asignaciones"] == asignaciones
    assert entorno.repo.filas[carga.id][0].extras["Región"] == "REGIÓN PRUEBA"
    assert entorno.repo.cargas[carga.id]["formato"].mapeo["ignoradas"]["4"] == "Región"


def test_procesar_dos_veces_la_misma_carga_confirmada_no_cambia_el_mapeo(entorno) -> None:
    primera = entorno.confirmar("datos.csv", CSV, cambios={3: None})
    segunda = entorno.confirmar("datos.csv", CSV, fecha=OTRA_FECHA, cambios={3: None})

    entorno.ingesta.procesar_carga(primera.id)
    entorno.ingesta.procesar_carga(segunda.id)

    assert (
        entorno.repo.cargas[primera.id]["formato"].mapeo["columnas"]
        == (entorno.repo.cargas[segunda.id]["formato"].mapeo["columnas"])
    )
