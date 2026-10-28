r"""Regenera los fixtures `tests/datos/sabana_ruta_*.xlsb` (T-SAB-1).

Son .xlsb SINTETICOS y pequenos: una hoja VENCIDA con unas pocas filas inventadas (telefonos
900000xxx) y el rango convertido en una Tabla de Excel (ListObject).

La causa del 400 del 26.10 NO es la Tabla: es la ruta de la carpeta desde la que se guardo el libro.
Excel la graba dentro de `xl/workbook.bin`, en el registro 2071 (BrtAbsPath15), como texto UTF-16
con un prefijo de 4 bytes (carga = 4 + 2 por caracter). La sabana del 26.10 la traia con 70
caracteres y python-calamine 0.8.2 daba 0 hojas («no tiene hojas»); desde la 0.9.0 se lee. Se
midio cambiando solo ese registro de un libro sano, con la 0.8.2: el largo del registro en 2 bytes
(62 caracteres o mas) NO basta, y de 10 a 200 caracteres solo fallan 70 (0 hojas) y 76 (panico de
Rust). Por eso hay dos fixtures:

- `sabana_ruta_larga.xlsb`: la carpeta grabada mide EXACTAMENTE 70 caracteres (falla con la 0.8.2);
- `sabana_ruta_corta.xlsb`: de control, con una de 30 (pasa con las dos).

PRIVACIDAD: la ruta queda grabada DENTRO del fixture versionado. Por eso se guarda siempre desde una
carpeta INVENTADA en la raiz del disco (no la del usuario, ni la de trabajo, ni una nube), que el
script crea y borra, y al final se comprueba que el .xlsb no contenga el usuario, el nombre de la
maquina ni rutas reales. Las carpetas son parametros (`--carpeta-larga`, `--carpeta-corta`).

Python no escribe .xlsb: se crea el .xlsx con openpyxl (con la tabla) y Excel (Windows) lo guarda
como .xlsb por COM, con SU PROPIA instancia (`New-Object -ComObject` crea otra): se cierra solo con
`Quit()`, nunca se mata un proceso EXCEL ni se toca un libro ajeno.

Uso, desde backend/:

    uv run python scripts/generar_fixture_xlsb_con_tabla.py
"""

import argparse
import getpass
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402

XL_EXCEL_12 = 50  # formato .xlsb en la API de Excel
DATOS = Path(__file__).resolve().parents[1] / "tests" / "datos"
CARPETA_LARGA = r"C:\sintetico_auto_cusco\carpeta_de_prueba_larga\subcarpeta_inventadas"
CARPETA_CORTA = r"C:\sintetico_auto_cusco\corta"
CABECERAS = ["Pagare", "Titular", "Teléfono", "Región", "Saldo Capital Pendiente"]
FILAS = [
    ["000000000000000001", "APELLIDO/NOMBRE FICTICIO", 900000001, "REGION PRUEBA", 1500.5],
    ["000000000000000002", "OTRO/FICTICIO", 900000002, "REGION PRUEBA", 320.0],
    ["000000000000000003", "TERCERO/FICTICIO", 900000003, "OTRA REGION", 78.25],
]
BRT_ABS_PATH_15 = 2071


def generar_xlsx(destino: Path) -> None:
    libro = openpyxl.Workbook()
    libro.properties.creator = libro.properties.lastModifiedBy = "sintetico"
    hoja = libro.active
    hoja.title = "VENCIDA"
    hoja.append(CABECERAS)
    for fila in FILAS:
        hoja.append(fila)
    referencia = f"A1:{openpyxl.utils.get_column_letter(len(CABECERAS))}{len(FILAS) + 1}"
    tabla = Table(displayName="TablaSabana", ref=referencia)
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    hoja.add_table(tabla)
    libro.save(destino)


def guardar_como_xlsb(origen: Path, destino: Path) -> None:
    """Excel abre el .xlsx y lo guarda como .xlsb en `destino`: esa carpeta queda grabada."""
    guion = f"""
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$excel.DisplayAlerts = $false
try {{
  $libro = $excel.Workbooks.Open('{origen}')
  # Excel graba el nombre de usuario y la empresa de Office: se pone un valor inventado.
  $enlace = [System.Reflection.BindingFlags]
  $propiedades = $libro.BuiltinDocumentProperties
  foreach ($nombre in 'Author', 'Last Author', 'Company', 'Manager') {{
    $com = [System.__ComObject]
    $p = $com.InvokeMember('Item', $enlace::GetProperty, $null, $propiedades, @($nombre))
    $com.InvokeMember('Value', $enlace::SetProperty, $null, $p, @('sintetico'))
  }}
  $libro.SaveAs('{destino}', {XL_EXCEL_12})
  $libro.Close($false)
}} finally {{
  $excel.Quit()
}}
"""
    subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", guion],
        check=True,
        capture_output=True,
    )


def registros(contenido: bytes):
    """Los registros BIFF12 de `xl/workbook.bin`: (tipo, carga, bytes que ocupa el largo)."""
    i = 0
    while i < len(contenido):
        tipo = contenido[i] & 0x7F
        if contenido[i] & 0x80:
            i += 1
            tipo |= (contenido[i] & 0x7F) << 7
        i += 1
        largo, bytes_del_largo = 0, 0
        while True:
            byte = contenido[i]
            i += 1
            largo |= (byte & 0x7F) << (7 * bytes_del_largo)
            bytes_del_largo += 1
            if not byte & 0x80:
                break
        yield tipo, contenido[i : i + largo], bytes_del_largo
        i += largo


def ruta_grabada(xlsb: Path) -> tuple[int, int]:
    """(largo de la carga, bytes del largo del registro) de BrtAbsPath15."""
    with zipfile.ZipFile(xlsb) as zip_:
        libro = zip_.read("xl/workbook.bin")
    for tipo, carga, bytes_del_largo in registros(libro):
        if tipo == BRT_ABS_PATH_15:
            return len(carga), bytes_del_largo
    raise AssertionError("el libro no trae BrtAbsPath15")


def anonimizar(xlsb: Path) -> set[str]:
    """Excel graba SIEMPRE el nombre de usuario de Office como `cp:lastModifiedBy`, aunque se pongan
    las propiedades del libro (el nombre sale de la configuracion de Office, que no se toca). Se
    reescribe el .xlsb con «sintetico» y se devuelve el nombre que tenia, para buscarlo luego."""
    encontrados: set[str] = set()
    destino = xlsb.with_suffix(".tmp")
    with zipfile.ZipFile(xlsb) as origen, zipfile.ZipFile(destino, "w") as nuevo:
        for item in origen.infolist():
            datos = origen.read(item.filename)
            if item.filename == "docProps/core.xml":
                texto = datos.decode("utf-8")
                for nombre in re.findall(r"<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", texto):
                    encontrados.add(nombre)
                texto = re.sub(
                    r"(<cp:lastModifiedBy>).*?(</cp:lastModifiedBy>)",
                    lambda m: m.group(1) + "sintetico" + m.group(2),
                    texto,
                )
                datos = texto.encode("utf-8")
            nuevo.writestr(item, datos)
    destino.replace(xlsb)
    return encontrados - {"sintetico"}


def comprobar_privacidad(xlsb: Path, ademas: set[str] = frozenset()) -> None:  # type: ignore[assignment]
    """Falla si algo de este equipo quedo grabado en el .xlsb (en ASCII o en UTF-16)."""
    prohibidas = {
        *ademas,
        os.environ.get("USERNAME", ""),
        getpass.getuser(),
        os.environ.get("COMPUTERNAME", ""),
        "Users",
        "Claude",
        "Google",
        "OneDrive",
        "AppData",
    } - {""}
    # Del usuario "NOMBRE 2" se busca cada palabra de 3 letras o mas.
    partes = [p for nombre in prohibidas for p in nombre.split() if len(p) >= 3]
    with zipfile.ZipFile(xlsb) as zip_:
        for nombre in zip_.namelist():
            datos = zip_.read(nombre).lower()
            if nombre.startswith("docProps/") and nombre.endswith(".xml"):
                assert b"@" not in datos, f"{nombre} contiene un correo"
            for parte in partes:
                for codificado in (parte.encode("ascii", "ignore"), parte.encode("utf-16-le")):
                    assert not codificado or codificado.lower() not in datos, (
                        f"{nombre} contiene {parte!r}: el fixture no es presentable"
                    )


def generar(destino: Path, carpeta: str, caracteres_grabados: int) -> None:
    with tempfile.TemporaryDirectory() as temporal:
        intermedio = Path(temporal) / "sabana.xlsx"
        generar_xlsx(intermedio)
        raiz = Path(carpeta)
        existia = raiz.exists()
        raiz.mkdir(parents=True, exist_ok=True)
        try:
            guardado = raiz / "sabana.xlsb"
            guardar_como_xlsb(intermedio, guardado)
            shutil.copyfile(guardado, destino)
        finally:
            # Solo se borra lo que creo este script; la raiz sintetica vacia tambien.
            if not existia:
                shutil.rmtree(Path(carpeta.split("\\")[0] + "\\" + carpeta.split("\\")[1]))
    largo_carga, bytes_largo = ruta_grabada(destino)
    assert (largo_carga - 4) // 2 == caracteres_grabados, (largo_carga, bytes_largo)
    comprobar_privacidad(destino, anonimizar(destino))
    with zipfile.ZipFile(destino) as zip_:
        tablas = [n for n in zip_.namelist() if n.startswith("xl/tables/")]
    print(
        f"{destino.name}: {destino.stat().st_size} bytes; carpeta de {len(carpeta)} caracteres; "
        f"BrtAbsPath15 con carga de {largo_carga} bytes ({(largo_carga - 4) // 2} caracteres) y "
        f"largo en {bytes_largo} byte(s); "
        f"tablas: {tablas}"
    )


def main() -> None:
    analizador = argparse.ArgumentParser(description=__doc__)
    analizador.add_argument("--carpeta-larga", default=CARPETA_LARGA)
    analizador.add_argument("--carpeta-corta", default=CARPETA_CORTA)
    opciones = analizador.parse_args()

    DATOS.mkdir(parents=True, exist_ok=True)
    generar(DATOS / "sabana_ruta_larga.xlsb", opciones.carpeta_larga, caracteres_grabados=70)
    generar(DATOS / "sabana_ruta_corta.xlsb", opciones.carpeta_corta, caracteres_grabados=30)


if __name__ == "__main__":
    main()
